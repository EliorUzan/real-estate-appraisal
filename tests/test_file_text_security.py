import zipfile

import pytest

from appraisal_assistant.services.file_text_extractor import FileTextExtractor


def test_oversized_attachment_rejected_before_parsing(tmp_path):
    path = tmp_path / "large.pdf"
    with path.open("wb") as stream:
        stream.truncate(20 * 1024 * 1024 + 1)
    with pytest.raises(ValueError, match="20MB"):
        FileTextExtractor().extract_text(path)


def test_high_expansion_office_document_rejected_before_parser(tmp_path):
    path = tmp_path / "large.docx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"A" * (100 * 1024 * 1024 + 1))
    with pytest.raises(ValueError, match="פריסה"):
        FileTextExtractor().extract_text(path)


def test_bounded_output_and_normal_hebrew_preview(tmp_path):
    path = tmp_path / "normal.txt"
    path.write_text("שלום עולם זה קובץ לבדיקה בלבד", encoding="utf-8")
    assert FileTextExtractor().first_words(path).words == "שלום עולם זה קובץ לבדיקה"
    with pytest.raises(ValueError, match="טקסט"):
        FileTextExtractor._bounded_join(["a" * 1_000_000, "b" * 1_000_000])


def test_tiny_sparse_spreadsheet_is_rejected_within_budget(tmp_path):
    import time
    from openpyxl import Workbook
    original = tmp_path / "normal.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "fixture"
    workbook.save(original)
    workbook.close()
    crafted = tmp_path / "sparse.xlsx"
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(crafted, "w", zipfile.ZIP_DEFLATED) as dest:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b'ref="A1:A1"', b'ref="A1:XFD1048576"').replace(b'<row r="1">', b'<row r="1048576">').replace(b'r="A1"', b'r="XFD1048576"')
            dest.writestr(entry, data)
    started = time.monotonic()
    with pytest.raises(ValueError, match="תאים"):
        FileTextExtractor().extract_text(crafted)
    assert time.monotonic() - started < 5


def test_regular_hebrew_office_documents_still_extract(tmp_path):
    from openpyxl import Workbook
    from docx import Document
    sheet = tmp_path / "regular.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "שלום"
    workbook.active["J100"] = "עולם"
    workbook.save(sheet)
    workbook.close()
    text = FileTextExtractor().extract_text(sheet)
    assert "שלום" in text and "עולם" in text
    doc = tmp_path / "regular.docx"
    document = Document()
    document.add_paragraph("שלום עולם")
    document.save(doc)
    assert "שלום עולם" in FileTextExtractor().extract_text(doc)


def test_extraction_deadline_becomes_actionable_error(tmp_path, monkeypatch):
    import subprocess
    def timeout(*args, **kwargs):
        assert kwargs["timeout"] == 20
        raise subprocess.TimeoutExpired("synthetic-parser", 20)
    monkeypatch.setattr(subprocess, "run", timeout)
    with pytest.raises(ValueError, match="זמן"):
        FileTextExtractor().extract_text(tmp_path / "fixture.pdf")
