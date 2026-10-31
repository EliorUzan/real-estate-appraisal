from __future__ import annotations

import csv
import json
import subprocess
import sys
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path


class ExtractionLimitError(ValueError):
    """Safe user-facing explanation for our explicit extraction limits."""


@dataclass(frozen=True)
class TextPreview:
    words: str
    warning: str | None = None


class FileTextExtractor:
    """Extracts a short local preview for development-time agent verification."""

    def first_words(self, path: Path, count: int = 5) -> TextPreview:
        try:
            text = self._read_text(path)
        except Exception as error:  # The agent must expose, not hide, an unreadable attachment.
            return TextPreview("לא ניתן לקרוא את הקובץ", f"לא ניתן לקרוא את {path.name}: {error}")
        words = re.findall(r"[\w\u0590-\u05FF]+(?:['׳\-][\w\u0590-\u05FF]+)*", text, flags=re.UNICODE)
        if not words:
            return TextPreview("לא נמצא טקסט קריא", f"לא נמצא טקסט קריא בקובץ {path.name}.")
        warning = None
        if path.suffix.lower() == ".doc":
            warning = f"הטקסט מתוך {path.name} חולץ מקובץ Word ישן ועלול להיות חלקי."
        return TextPreview(" ".join(words[:count]), warning)

    def extract_text(self, path: Path) -> str:
        """Return full readable text for a provider that cannot accept the raw format."""
        return self._read_text(path)

    def _read_text(self, path: Path) -> str:
        """Run untrusted parsers outside the UI/provider process with a hard deadline."""
        path = path.resolve()
        worker = Path(__file__).with_name("extraction_worker.py")
        try:
            result = subprocess.run(
                [sys.executable, "-I", str(worker), str(path)],
                capture_output=True, timeout=20,
                creationflags=subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0,
            )
        except subprocess.TimeoutExpired as error:
            raise ValueError("עיבוד הקובץ חרג ממגבלת הזמן. יש לפצל או להמיר אותו.") from error
        if result.returncode or len(result.stdout) > 12_000_000:
            raise ValueError("לא ניתן לעבד את הקובץ במסגרת מגבלות הבטיחות")
        payload = json.loads(result.stdout)
        if "error" in payload:
            raise ValueError(payload["error"])
        return payload["text"]

    def _read_text_in_worker(self, path: Path) -> str:
        if not path.is_file() or path.stat().st_size > 20 * 1024 * 1024:
            raise ExtractionLimitError("גודל הקובץ המרבי הוא 20MB")
        if path.suffix.lower() in {".docx", ".xlsx"}:
            with zipfile.ZipFile(path) as archive:
                entries = archive.infolist()
                if len(entries) > 10000 or sum(entry.file_size for entry in entries) > 100 * 1024 * 1024:
                    raise ExtractionLimitError("תוכן הקובץ לאחר פריסה גדול מדי")
        text = self._read_supported_text(path)
        if len(text) > 2_000_000:
            raise ExtractionLimitError("כמות הטקסט בקובץ גדולה מדי")
        return text

    @staticmethod
    def _bounded_join(values) -> str:
        parts, length = [], 0
        for value in values:
            length += len(value) + 1
            if length > 2_000_000:
                raise ExtractionLimitError("כמות הטקסט בקובץ גדולה מדי")
            parts.append(value)
        return "\n".join(parts)

    def _read_supported_text(self, path: Path) -> str:
        suffix = path.suffix.lower()
        if suffix in {".txt", ".md"}:
            return self._read_plain_text(path)
        if suffix == ".csv":
            return self._read_csv(path)
        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(path)
            if len(reader.pages) > 2000:
                raise ExtractionLimitError("מספר העמודים בקובץ גדול מדי")
            return self._bounded_join(page.extract_text() or "" for page in reader.pages)
        if suffix == ".docx":
            from docx import Document

            document = Document(path)
            return self._bounded_join(paragraph.text for paragraph in document.paragraphs)
        if suffix == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                if len(workbook.worksheets) > 100:
                    raise ExtractionLimitError("מספר הגיליונות גדול מדי")
                visited = 0
                for sheet in workbook.worksheets:
                    visited += (sheet.max_row or 0) * (sheet.max_column or 0)
                    if visited > 1_000_000:
                        raise ExtractionLimitError("מספר התאים בקובץ גדול מדי; יש לצמצם את טווח הגיליון")
                return self._bounded_join(str(value) for sheet in workbook.worksheets for row in sheet.iter_rows(values_only=True) for value in row if value is not None)
            finally:
                workbook.close()
        if suffix == ".xls":
            import xlrd

            workbook = xlrd.open_workbook(path)
            try:
                return self._bounded_join(str(sheet.cell_value(row, column)) for sheet in workbook.sheets() for row in range(sheet.nrows) for column in range(sheet.ncols))
            finally:
                workbook.release_resources()
        if suffix == ".doc":
            return self._read_legacy_doc(path)
        raise ExtractionLimitError("סוג קובץ אינו נתמך")

    @staticmethod
    def _read_plain_text(path: Path) -> str:
        for encoding in ("utf-8-sig", "utf-8", "cp1255"):
            try:
                return path.read_text(encoding=encoding)
            except UnicodeDecodeError:
                continue
        return path.read_text(encoding="latin-1")

    def _read_csv(self, path: Path) -> str:
        text = self._read_plain_text(path)
        return "\n".join(" ".join(row) for row in csv.reader(text.splitlines()))

    @staticmethod
    def _read_legacy_doc(path: Path) -> str:
        """Best-effort fallback for legacy binary Word files; production should use a dedicated converter."""
        raw = path.read_bytes()
        candidates = [raw.decode(encoding, errors="ignore") for encoding in ("utf-16le", "cp1255", "latin-1")]
        return max(candidates, key=lambda value: len(re.findall(r"[\w\u0590-\u05FF]+", value)))
